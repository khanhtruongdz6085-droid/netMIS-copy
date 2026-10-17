from pathlib import Path
import subprocess,os,io
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT=Path(__file__).resolve().parent
IMG=ROOT/'images'
CACHE=ROOT/'.cache';CACHE.mkdir(exist_ok=True)
RUN_ENV=dict(os.environ,XDG_CACHE_HOME=str(CACHE))
DATA_DIR=ROOT/'pandoc-data'
(DATA_DIR/'translations').mkdir(parents=True,exist_ok=True)
(DATA_DIR/'translations/vi.yaml').write_text('''Abstract: Tóm tắt
Appendix: Phụ lục
Bibliography: Tài liệu tham khảo
Cc: Cùng gửi
Chapter: Chương
Contents: Mục lục
Encl: Kèm theo
Figure: Hình
Glossary: Thuật ngữ
Index: Chỉ mục
Listing: Đoạn mã
ListOfFigures: Danh mục hình
ListOfTables: Danh mục bảng
Page: Trang
Part: Phần
Preface: Lời nói đầu
Proof: Chứng minh
References: Tài liệu tham khảo
See: Xem
SeeAlso: Xem thêm
Table: Bảng
To: Gửi
''',encoding='utf-8')
architecture=r'''digraph G {
 graph [rankdir=TB, bgcolor="white", fontname="DejaVu Sans", nodesep=0.35, ranksep=0.55, pad=0.2];
 node [shape=box, style="rounded,filled", fillcolor="#eef5fb", color="#2a5575", fontname="DejaVu Sans", fontsize=12, margin="0.15,0.10"];
 edge [color="#54758a", fontname="DejaVu Sans", fontsize=10, arrowsize=0.7];
 UI [label="KHỞI TẠO VÀ GIAO DIỆN\nMain • SampleData\nConsoleMenu • ConsoleInput • ReportMenu",fillcolor="#dceaf5"];
 Facade [label="NetCafe\nĐiểm gọi chung của các nghiệp vụ"];
 Services [label="NGHIỆP VỤ\nCatalogService • SessionService\nFoodOrderService • ReceiptService • PaymentService",fillcolor="#e9f4eb"];
 Models [label="MÔ HÌNH\nKhu • Máy • Khách • Nhân viên • Món\nPhiên • Đơn • Chi tiết • Hóa đơn • Thanh toán"];
 DAOs [label="TRUY CẬP DỮ LIỆU\n9 DAO kế thừa MemoryDao<T>\nMỗi DAO giữ một List<T>"];
 Context [label="NetCafeData\nTạo và giữ 9 DAO\nThực hiện DaoQueryContext",fillcolor="#fff2db"];
 Contracts [label="HỢP ĐỒNG VÀ HỖ TRỢ\nIdentifiable • Chargeable\nChecks • ServiceChecks • DaoQuerySupport\nEnum • Exception",fillcolor="#f1eef7"];
 UI -> Facade [label="gọi thao tác"];
 Facade -> Services;
 UI -> DAOs [label="xem báo cáo",constraint=false];
 Services -> Models [label="kiểm tra, thay đổi trạng thái"];
 Services -> DAOs [label="kiểm tra, lưu"];
 DAOs -> Context [label="đọc danh sách liên quan",style=dashed];
 Context -> DAOs [label="giữ các DAO",style=dashed,constraint=false];
 Models -> Contracts [style=dashed,label="sử dụng / thực hiện"];
 DAOs -> Contracts [style=dashed,constraint=false];
 {rank=same;Models;DAOs;}
}'''
relationships=r'''digraph G {
 graph [rankdir=TB,bgcolor="white",fontname="DejaVu Sans",nodesep=0.65,ranksep=1.0,pad=0.25,splines=polyline];
 node [shape=box,style="rounded,filled",fillcolor="#edf5fb",color="#2a5575",fontname="DejaVu Sans",fontsize=12,margin="0.15,0.12"];
 edge [color="#54758a",fontname="DejaVu Sans",fontsize=10,arrowsize=0.65,labeldistance=1.8];
 ComputerSection [label="ComputerSection\nKhu máy"];
 Computer [label="Computer\nMáy tính"];
 Customer [label="Customer\nKhách hàng"];
 Employee [label="Employee\nNhân viên"];
 Session [label="Session\nPhiên chơi",fillcolor="#e9f4eb"];
 FoodService [label="FoodService\nLần gọi món"];
 Food [label="Food\nMón / tồn kho"];
 Receipt [label="Receipt\nHóa đơn",fillcolor="#fff2db"];
 FoodOrderDetail [label="FoodOrderDetail\nDòng món"];
 Payment [label="Payment\nThanh toán",fillcolor="#fff2db"];
 ComputerSection -> Computer [label="chứa",taillabel="1",headlabel="0..n"];
 Computer -> Session [label="được dùng",taillabel="1",headlabel="0..n"];
 Customer -> Session [label="chơi",taillabel="1",headlabel="0..n"];
 Employee -> Session [label="mở phiên\n1 → 0..n"];
 Session -> FoodService [label="có lần gọi",taillabel="1",headlabel="0..n"];
 Employee -> FoodService [label="xử lý\n1 → 0..n",constraint=false];
 FoodService -> FoodOrderDetail [label="gồm các dòng",taillabel="1",headlabel="1..n"];
 Food -> FoodOrderDetail [label="được gọi",taillabel="1",headlabel="0..n"];
 Session -> Receipt [label="được lập",taillabel="1",headlabel="0..1"];
 Employee -> Receipt [label="lập\n1 → 0..n",constraint=false];
 Receipt -> Payment [label="được thanh toán",taillabel="1",headlabel="0..1"];
 {rank=same;Computer;Customer;Employee;}
 {rank=same;FoodService;Food;Receipt;}
 {rank=same;FoodOrderDetail;Payment;}
}'''

def diagram(name,data):
    path=IMG/(name+'.dot');path.write_text(data,encoding='utf-8')
    subprocess.run(['dot','-Tpng','-Gdpi=180',str(path),'-o',str(IMG/(name+'.png'))],check=True,env=RUN_ENV)
    subprocess.run(['dot','-Tsvg',str(path),'-o',str(IMG/(name+'.svg'))],check=True,env=RUN_ENV)

diagram('kien-truc',architecture);diagram('quan-he',relationships)
default_reference=subprocess.run(['pandoc','--print-default-data-file','reference.docx'],check=True,capture_output=True).stdout
doc=Document(io.BytesIO(default_reference));sec=doc.sections[0]
for style in doc.styles:
    if style.style_id.startswith('Heading') and style.style_id[7:].isdigit():
        style.element.find(qn('w:name')).set(qn('w:val'),'heading '+style.style_id[7:])
sec.page_height=Cm(29.7);sec.page_width=Cm(21)
sec.left_margin=Cm(2.3);sec.right_margin=Cm(2.0);sec.top_margin=Cm(2.1);sec.bottom_margin=Cm(2.0)
sec.header_distance=Cm(0.8);sec.footer_distance=Cm(0.8)
normal=doc.styles['Normal'];normal.font.name='Times New Roman';normal.font.size=Pt(12)
normal.paragraph_format.line_spacing=1.12;normal.paragraph_format.space_after=Pt(6)
for name,size in [('Title',24),('Subtitle',13),('Heading 1',16),('Heading 2',13),('Heading 3',12)]:
    st=doc.styles[name];st.font.name='Times New Roman';st.font.size=Pt(size)
    if name.startswith('Heading'):
        st.font.color.rgb=RGBColor.from_string('19486A');st.font.bold=True
        st.paragraph_format.keep_with_next=True;st.paragraph_format.space_before=Pt(12);st.paragraph_format.space_after=Pt(6)
for name in ['Source Code']:
    if name in doc.styles:
        st=doc.styles[name];st.font.name='DejaVu Sans Mono';st.font.size=Pt(9)
header=sec.header.paragraphs[0];header.text='NETCAFEMIS  |  ĐẶC TẢ DỰ ÁN'
header.runs[0].font.size=Pt(9);header.runs[0].font.color.rgb=RGBColor.from_string('667788')
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('Trang ').font.size=Pt(9)
for instr in ['PAGE','NUMPAGES']:
    if instr=='NUMPAGES':footer.add_run(' / ').font.size=Pt(9)
    r=footer.add_run();fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),instr);r._r.addnext(fld)
doc.save(ROOT/'reference.docx')
output=ROOT/'Dac-ta-NetCafeMIS.docx'
subprocess.run(['pandoc',str(ROOT/'DAC_TA_NETCAFEMIS.md'),'--from=markdown','--to=docx','--standalone','--data-dir='+str(DATA_DIR),'--resource-path='+str(ROOT),'--reference-doc='+str(ROOT/'reference.docx'),'-o',str(output)],check=True,env=RUN_ENV)
doc=Document(output)
first=next(p for p in doc.paragraphs if p.style.name=='Heading 1')
chapters=[p.text for p in doc.paragraphs if p.style.name=='Heading 1']
toc=first.insert_paragraph_before('MỤC LỤC NỘI DUNG',style='Heading 1');toc.paragraph_format.page_break_before=True
for title in chapters:
    p=first.insert_paragraph_before(title);p.paragraph_format.space_after=Pt(4)
    for run in p.runs:run.font.size=Pt(11)
first.paragraph_format.page_break_before=True
for p in doc.paragraphs:
    if p.style.name in ('Title','Subtitle','Date'):
        p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    if p.style.name=='Title':p.paragraph_format.space_before=Pt(100);p.paragraph_format.space_after=Pt(18)
    if p.style.name=='Caption':
        p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        for run in p.runs:run.font.size=Pt(10)
    if p.style.name=='Heading 1' and p.text.startswith(('4. ','7. ','9. ','14. ')):
        p.paragraph_format.page_break_before=True
for table in doc.tables:
    table.style='Table'
    borders=OxmlElement('w:tblBorders')
    for side in ['top','left','bottom','right','insideH','insideV']:
        edge=OxmlElement('w:'+side);edge.set(qn('w:val'),'single');edge.set(qn('w:sz'),'4');edge.set(qn('w:color'),'9DAEBB');borders.append(edge)
    table._tbl.tblPr.append(borders)
    table.autofit=True
    for i,row in enumerate(table.rows):
        trPr=row._tr.get_or_add_trPr();cant=OxmlElement('w:cantSplit');trPr.append(cant)
        if i==0 and trPr.find(qn('w:tblHeader')) is None:
            repeat=OxmlElement('w:tblHeader');trPr.append(repeat)
        for cell in row.cells:
            if i==0:
                shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E6EFF6');cell._tc.get_or_add_tcPr().append(shade)
            for p in cell.paragraphs:
                p.paragraph_format.space_after=Pt(3);p.paragraph_format.space_before=Pt(2);p.paragraph_format.line_spacing=1.0
                p.paragraph_format.keep_with_next=False;p.paragraph_format.keep_together=False
                for run in p.runs:
                    run.font.name='Times New Roman';run.font.size=Pt(10.5)
                    if i==0:run.bold=True
settings=doc.settings.element
update=OxmlElement('w:updateFields');update.set(qn('w:val'),'true');settings.append(update)
doc.core_properties.title='Đặc tả dự án NetCafeMIS'
doc.core_properties.subject='Java OOP, DAO trong bộ nhớ, phiên chơi, gọi món, hóa đơn và 30 truy vấn'
doc.core_properties.author='NetCafeMIS'
doc.save(output)
print('Saved',output,'paragraphs=',len(doc.paragraphs),'tables=',len(doc.tables),'diagrams=',len(doc.inline_shapes))
